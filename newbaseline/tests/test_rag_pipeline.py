"""Focused offline checks for the paper-compatible RAG flow."""

from __future__ import annotations

import json
import hashlib
import tempfile
import unittest
from pathlib import Path

import numpy as np
from docx import Document

from newbaseline.src.rag.anchor_hierarchy import AnchorHierarchy, build_document_description
from newbaseline.src.rag.corpus import PaperEmbeddingCorpus
from newbaseline.src.rag.lexical import lexical_query
from newbaseline.src.rag.router import PAPER_SERIES, PaperNNRouter
from newbaseline.src.rag.service import PaperRagService
from newbaseline.src.rag.types import RetrievalHit
from newbaseline.src.rag.vocabulary import Vocabulary


class PaperEmbeddingCorpusTests(unittest.TestCase):
    def test_retrieval_uses_metadata_source_index_and_retains_empty_router_label(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chunk_root = root / "chunks"
            chunk_root.mkdir()
            (chunk_root / "ChunkSeries21.json").write_text(
                json.dumps([
                    {"chunk_id": "original-0", "text": "less relevant"},
                    {"chunk_id": "original-1", "text": "most relevant"},
                ]),
                encoding="utf-8",
            )
            np.save(root / "series21.npy", np.asarray([[0.0, 1.0], [1.0, 0.0]], dtype=np.float32))
            (root / "series21.jsonl").write_text(
                "\n".join(
                    json.dumps(item)
                    for item in (
                        {"chunk_id": "original-0", "source_chunk_index": 0, "document_name": "a"},
                        {"chunk_id": "original-1", "source_chunk_index": 1, "document_name": "b"},
                    )
                )
                + "\n",
                encoding="utf-8",
            )
            (root / "manifest.json").write_text(
                json.dumps(
                    {
                        "source_chunk_directory": str(chunk_root),
                        "series": {
                            "21": {
                                "vector_file": "series21.npy",
                                "metadata_file": "series21.jsonl",
                                "chunk_file": "ChunkSeries21.json",
                            }
                        },
                    }
                ),
                encoding="utf-8",
            )
            corpus = PaperEmbeddingCorpus(root)
            hits, searched, empty = corpus.search(["25", "21"], np.asarray([1.0, 0.0]), top_k=1)
        self.assertEqual(empty, ["25"])
        self.assertEqual(searched, ["21"])
        self.assertEqual(hits[0].text, "most relevant")
        self.assertEqual(hits[0].metadata["chunk_id"], "original-1")

    def test_search_many_preserves_query_order(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chunk_root = root / "chunks"
            chunk_root.mkdir()
            (chunk_root / "ChunkSeries21.json").write_text(
                json.dumps([
                    {"chunk_id": "x", "text": "x"},
                    {"chunk_id": "y", "text": "y"},
                ]),
                encoding="utf-8",
            )
            np.save(root / "series21.npy", np.asarray([[1.0, 0.0], [0.0, 1.0]], dtype=np.float32))
            (root / "series21.jsonl").write_text(
                "\n".join(json.dumps({"chunk_id": item, "source_chunk_index": index}) for index, item in enumerate(("x", "y")))
                + "\n",
                encoding="utf-8",
            )
            (root / "manifest.json").write_text(
                json.dumps({"source_chunk_directory": str(chunk_root), "series": {"21": {"vector_file": "series21.npy", "metadata_file": "series21.jsonl", "chunk_file": "ChunkSeries21.json"}}}),
                encoding="utf-8",
            )
            corpus = PaperEmbeddingCorpus(root)
            hits_by_query, searched, empty = corpus.search_many(
                ["21"], np.asarray([[1.0, 0.0], [0.0, 1.0]], dtype=np.float32), top_k=1
            )
        self.assertEqual(searched, ["21"])
        self.assertEqual(empty, [])
        self.assertEqual([hits[0].metadata["chunk_id"] for hits in hits_by_query], ["x", "y"])

    def test_hybrid_search_fuses_dense_and_lexical_rankings(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chunk_root = root / "chunks"
            chunk_root.mkdir()
            chunks = [
                {
                    "chunk_id": f"chunk-{index:02d}",
                    "heading": "Procedure",
                    "text": "rareprotocol exact procedure" if index == 9 else f"generic network text {index}",
                }
                for index in range(10)
            ]
            (chunk_root / "ChunkSeries21.json").write_text(json.dumps(chunks), encoding="utf-8")
            np.save(
                root / "series21.npy",
                np.asarray([[1.0 - index * 0.05, 0.0] for index in range(10)], dtype=np.float32),
            )
            (root / "series21.jsonl").write_text(
                "\n".join(
                    json.dumps(
                        {
                            "chunk_id": chunk["chunk_id"],
                            "source_chunk_index": index,
                            "heading": chunk["heading"],
                        }
                    )
                    for index, chunk in enumerate(chunks)
                )
                + "\n",
                encoding="utf-8",
            )
            (root / "manifest.json").write_text(
                json.dumps(
                    {
                        "source_chunk_directory": str(chunk_root),
                        "series": {
                            "21": {
                                "vector_file": "series21.npy",
                                "metadata_file": "series21.jsonl",
                                "chunk_file": "ChunkSeries21.json",
                            }
                        },
                    }
                ),
                encoding="utf-8",
            )
            PaperEmbeddingCorpus(root).build_lexical_index(root / "lexical.sqlite3")
            corpus = PaperEmbeddingCorpus(root, retrieval_backend="hybrid")
            hits, searched, empty = corpus.search(
                ["21"],
                np.asarray([1.0, 0.0]),
                top_k=2,
                query_text="rareprotocol",
            )

        by_id = {hit.metadata["chunk_id"]: hit for hit in hits}
        self.assertEqual(searched, ["21"])
        self.assertEqual(empty, [])
        self.assertEqual(set(by_id), {"chunk-00", "chunk-09"})
        self.assertEqual(by_id["chunk-00"].dense_rank, 1)
        self.assertEqual(by_id["chunk-09"].lexical_rank, 1)
        self.assertIsNone(by_id["chunk-09"].dense_rank)
        self.assertTrue(all(hit.retrieval_method == "hybrid" for hit in hits))

    def test_weighted_rrf_can_favor_dense_or_bm25_ranks(self) -> None:
        dense = [
            RetrievalHit(1.0, "21", "dense", {"chunk_id": "dense"}),
            *[
                RetrievalHit(0.9, "21", f"filler-{index}", {"chunk_id": f"filler-{index}"})
                for index in range(49)
            ],
            RetrievalHit(0.1, "21", "lexical", {"chunk_id": "lexical"}),
        ]
        lexical = [RetrievalHit(1.0, "21", "lexical", {"chunk_id": "lexical"})]

        dense_heavy = PaperEmbeddingCorpus._fuse_ranked_hits(
            dense,
            lexical,
            top_k=1,
            dense_weight=0.7,
            bm25_weight=0.3,
        )
        bm25_heavy = PaperEmbeddingCorpus._fuse_ranked_hits(
            dense,
            lexical,
            top_k=1,
            dense_weight=0.3,
            bm25_weight=0.7,
        )

        self.assertEqual(dense_heavy[0].metadata["chunk_id"], "dense")
        self.assertEqual(bm25_heavy[0].metadata["chunk_id"], "lexical")
        with self.assertRaises(ValueError):
            PaperEmbeddingCorpus._fuse_ranked_hits(
                dense,
                lexical,
                top_k=1,
                dense_weight=0.7,
                bm25_weight=0.4,
            )

    def test_hierarchical_search_ranks_all_series_then_keeps_rrf(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chunk_root = root / "chunks"
            chunk_root.mkdir()
            for series, document_key, chunk_id, text in (
                ("21", "21_series/21001", "s21", "normal chunk"),
                ("22", "22_series/22001", "s22", "lexical token"),
            ):
                (chunk_root / f"ChunkSeries{series}.json").write_text(
                    json.dumps([{"chunk_id": chunk_id, "text": text}]), encoding="utf-8"
                )
                np.save(root / f"series{series}.npy", np.asarray([[1.0, 0.0]], dtype=np.float32))
                (root / f"series{series}.jsonl").write_text(
                    json.dumps(
                        {
                            "chunk_id": chunk_id,
                            "source_chunk_index": 0,
                            "document_key": document_key,
                        }
                    )
                    + "\n",
                    encoding="utf-8",
                )
            (root / "manifest.json").write_text(
                json.dumps(
                    {
                        "source_chunk_directory": str(chunk_root),
                        "series": {
                            "21": {"vector_file": "series21.npy", "metadata_file": "series21.jsonl", "chunk_file": "ChunkSeries21.json"},
                            "22": {"vector_file": "series22.npy", "metadata_file": "series22.jsonl", "chunk_file": "ChunkSeries22.json"},
                        },
                    }
                ),
                encoding="utf-8",
            )
            corpus = PaperEmbeddingCorpus(root)
            hierarchy = AnchorHierarchy(
                series_vectors={"21": np.asarray([1.0, 0.0]), "22": np.asarray([0.0, 1.0])},
                document_vectors={
                    "21_series/21001": np.asarray([0.0, 1.0]),
                    "22_series/22001": np.asarray([1.0, 0.0]),
                },
                provenance={"artifact_version": 1},
            )
            hits, searched = corpus.search_hierarchical(
                np.asarray([1.0, 0.0]),
                1,
                query_text=None,
                hierarchy=hierarchy,
                series_weight=0.1,
                document_weight=0.1,
                chunk_weight=0.8,
            )
        self.assertEqual(searched, ["21", "22"])
        self.assertEqual(hits[0].metadata["chunk_id"], "s21")
        self.assertAlmostEqual(hits[0].anchor_series_score or 0.0, 1.0)
        self.assertAlmostEqual(hits[0].anchor_document_score or 0.0, 0.0)
        self.assertAlmostEqual(hits[0].anchor_chunk_score or 0.0, 1.0)
        self.assertAlmostEqual(hits[0].anchor_hierarchical_score or 0.0, 0.9)

    def test_document_description_uses_title_and_scope_with_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            raw_path = Path(temp_dir) / "raw.md"
            raw_path.write_text(
                "# 3GPP TS 21.001 V18.0.0\n## **System title**\n## --- 1 Scope\nUseful scope text.\n## --- 2 References\n",
                encoding="utf-8",
            )
            metadata = {
                "document_type": "TS",
                "document_number": "21.001",
                "series": "21",
                "headings": [
                    {"heading": "1", "title": "Scope", "line": 3},
                    {"heading": "2", "title": "References", "line": 5},
                ],
            }
            description = build_document_description(metadata, raw_path)
        self.assertIn("Title: System title.", description)
        self.assertIn("Scope: Useful scope text.", description)
        self.assertEqual(
            build_document_description({"document_type": "TR", "document_number": "22.001", "series": "22"}, None),
            "3GPP TR 22.001, series 22.",
        )

    def test_document_description_recovers_malformed_scope_heading_and_docx_summary(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            markdown_path = root / "raw.md"
            markdown_path.write_text(
                "# 1527 1 Scope\nApplies to system work.\n# 2 References\nIgnored.", encoding="utf-8"
            )
            markdown_description = build_document_description(
                {"document_type": "TS", "document_number": "29.002", "series": "29", "headings": []}, markdown_path
            )
            docx_path = root / "rel_17.docx"
            document = Document()
            document.add_paragraph("3GPP TR 21.915 V17.0.0")
            document.add_paragraph("Release 17 Description; Summary of Rel-17 Work Items")
            document.add_paragraph("1 Scope")
            document.add_paragraph("This summary covers Release 17 work items.")
            document.add_paragraph("2 References")
            document.add_paragraph("Ignored references.")
            document.save(docx_path)
            docx_description = build_document_description(
                {"document_type": "TR", "document_number": "21.915", "series": "release-summaries"}, docx_path
            )
        self.assertIn("Scope: 1527 1 Scope Applies to system work.", markdown_description)
        self.assertNotIn("Ignored.", markdown_description)
        self.assertIn("Title: Release 17 Description; Summary of Rel-17 Work Items.", docx_description)
        self.assertIn("Scope: This summary covers Release 17 work items.", docx_description)
        self.assertNotIn("Ignored references.", docx_description)

    def test_hierarchy_artifact_validation_and_hybrid_trace_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            corpus_manifest = root / "manifest.json"
            corpus_manifest.write_text("{}", encoding="utf-8")
            descriptions = root / "descriptions.jsonl"
            descriptions.write_text('{"kind":"series","key":"21","text":"series"}\n', encoding="utf-8")
            np.savez_compressed(
                root / "vectors.npz",
                series_keys=np.asarray(["21"]),
                series_vectors=np.asarray([[1.0, 0.0]], dtype=np.float32),
                document_keys=np.asarray(["21_series/21001"]),
                document_vectors=np.asarray([[0.0, 1.0]], dtype=np.float32),
            )
            manifest = {
                "artifact_version": 1,
                "corpus_manifest_sha256": hashlib.sha256(corpus_manifest.read_bytes()).hexdigest(),
                "embedding_backend": "test",
                "embedding_model": "test-model",
                "embedding_dimensions": 2,
                "vectors_file": "vectors.npz",
                "descriptions_file": "descriptions.jsonl",
                "descriptions_sha256": hashlib.sha256(descriptions.read_bytes()).hexdigest(),
            }
            (root / "hierarchy.json").write_text(json.dumps(manifest), encoding="utf-8")
            hierarchy = AnchorHierarchy.load(
                root,
                manifest_file="hierarchy.json",
                vectors_file="vectors.npz",
                corpus_manifest_sha256=manifest["corpus_manifest_sha256"],
                embedding_backend="test",
                embedding_model="test-model",
                dimensions=2,
            )
            self.assertEqual(set(hierarchy.document_vectors), {"21_series/21001"})
            with self.assertRaises(ValueError):
                AnchorHierarchy.load(
                    root,
                    manifest_file="hierarchy.json",
                    vectors_file="vectors.npz",
                    corpus_manifest_sha256="wrong",
                    embedding_backend="test",
                    embedding_model="test-model",
                    dimensions=2,
                )

        dense = RetrievalHit(
            0.9,
            "21",
            "dense",
            {"chunk_id": "dense"},
            anchor_series_score=0.2,
            anchor_document_score=0.3,
            anchor_chunk_score=0.9,
            anchor_hierarchical_score=0.77,
        )
        fused = PaperEmbeddingCorpus._fuse_ranked_hits([dense], [], top_k=1)[0]
        self.assertEqual(fused.retrieval_method, "hybrid")
        self.assertEqual(fused.anchor_hierarchical_score, 0.77)

    def test_lexical_query_keeps_technical_terms_and_drops_question_words(self) -> None:
        self.assertEqual(
            lexical_query("What is AMF on the N2 interface? [3GPP Release 18]"),
            '"amf" OR "n2" OR "interface"',
        )


class PaperRagServiceTests(unittest.TestCase):
    def test_service_exposes_selected_empty_and_searched_series(self) -> None:
        events: list[str] = []

        class FakeClient:
            def rephrase(self, question: str) -> str:
                events.append("rephrase")
                return f"rephrased {question}"

            def extract_facets(self, question: str) -> list[str]:
                events.append("facets")
                return ["meaning of AMF"]

            def embed(self, text: str) -> list[float]:
                events.append(f"query_embedding:{text}")
                return [1.0, 0.0]

            def embed_many(self, texts: list[str]) -> list[list[float]]:
                events.append(f"facet_embeddings:{'|'.join(texts)}")
                return [[1.0, 0.0] for _ in texts]

            def answer(self, question: str, contexts: list[str]) -> str:
                return "answer"

        class FakeRouter:
            def route(self, query_embedding: np.ndarray, top_k: int) -> list[str]:
                events.append("route")
                return ["25", "21"]

        class FakeCorpus:
            def search(
                self,
                selected: list[str],
                embedding: np.ndarray,
                top_k: int,
                query_text: str | None = None,
            ):
                events.append("search")
                hit = RetrievalHit(0.8, "21", "retrieved text", {"chunk_id": "original-1"})
                return [hit], ["21", "release-summaries"], ["25"]

            def expand_citations(self, seeds, **kwargs):
                return seeds, []

        service = object.__new__(PaperRagService)
        service.client = FakeClient()
        service.router = FakeRouter()
        service.corpus = FakeCorpus()
        service.vocabulary = Vocabulary({}, {})
        service.anchor_strategy = "router"
        service.anchor_hierarchy = None
        settings = {
            "anchor_strategy": "router",
            "router_top_k": 5,
            "retrieval_top_k": 8,
            "citation_strategy": "gain",
            "citation_max_depth": 1,
            "citation_max_chunks": 4,
            "citation_chunks_per_heading": 1,
            "citation_min_gain": 0.01,
        }
        service.settings = type("Settings", (), {"get": lambda _, _section, key: settings[key]})()
        result = service.run("What is AMF?", include_answer=False)
        self.assertEqual(result.router_selected_series, ["25", "21"])
        self.assertEqual(result.empty_selected_series, ["25"])
        self.assertEqual(result.searched_series, ["21", "release-summaries"])
        self.assertEqual(result.query_facets, ["meaning of AMF"])
        self.assertEqual(events[0], "rephrase")
        self.assertTrue(events[1].startswith("query_embedding:rephrased What is AMF?"))
        self.assertEqual(events[2:6], ["route", "search", "facets", "facet_embeddings:meaning of AMF"])
        self.assertIsNone(result.answer)

    def test_hierarchical_service_skips_router(self) -> None:
        events: list[str] = []

        class FakeClient:
            def rephrase(self, question: str) -> str:
                return question

            def extract_facets(self, question: str) -> list[str]:
                return [question]

            def embed(self, text: str) -> list[float]:
                return [1.0, 0.0]

            def embed_many(self, texts: list[str]) -> list[list[float]]:
                return [[1.0, 0.0] for _ in texts]

            def answer(self, question: str, contexts: list[str]) -> str:
                return "answer"

        class ExplodingRouter:
            def route(self, *_args):
                raise AssertionError("hierarchical mode must not call router")

        class FakeCorpus:
            def search_hierarchical(self, *_args, **_kwargs):
                events.append("hierarchical")
                return [RetrievalHit(0.9, "21", "text", {"chunk_id": "seed"})], ["21", "release-summaries"]

            def expand_citations(self, seeds, **kwargs):
                return seeds, []

        service = object.__new__(PaperRagService)
        service.client = FakeClient()
        service.router = ExplodingRouter()
        service.corpus = FakeCorpus()
        service.vocabulary = Vocabulary({}, {})
        service.anchor_strategy = "hierarchical"
        service.anchor_hierarchy = type("Hierarchy", (), {"provenance": {"artifact_version": 1}})()
        values = {
            "retrieval_top_k": 8,
            "anchor_series_weight": 0.1,
            "anchor_document_weight": 0.1,
            "anchor_chunk_weight": 0.8,
            "citation_strategy": "semantic_bfs",
            "citation_max_depth": 1,
            "citation_max_chunks": 4,
            "citation_chunks_per_heading": 1,
            "citation_min_gain": 0.01,
        }
        service.settings = type("Settings", (), {"get": lambda _, _section, key: values[key]})()

        result = service.run("What is AMF?", include_answer=False)

        self.assertEqual(events, ["hierarchical"])
        self.assertEqual(result.router_selected_series, [])
        self.assertEqual(result.searched_series, ["21", "release-summaries"])
        self.assertEqual(result.anchor_strategy, "hierarchical")

    def test_answer_context_preserves_seed_format_and_appends_neutral_citation(self) -> None:
        seed = RetrievalHit(
            0.8,
            "21",
            "seed text",
            {"chunk_id": "seed-id", "document_id": "21001", "heading": "1"},
        )
        citation = RetrievalHit(
            0.2,
            "23",
            "citation text",
            {"chunk_id": "cite-id", "document_id": "23501", "heading": "4"},
            origin="citation",
            citation_depth=1,
            parent_chunk_id="seed-id",
            retrieval_method="citation_gain",
            citation_gain=0.2,
            citation_facets=("the condition required for registration",),
        )

        contexts = PaperRagService._format_contexts([citation, seed])

        self.assertEqual(
            contexts[0],
            "[series=21; document=21001; heading=1; chunk_id=seed-id]\nseed text",
        )
        self.assertIn("Referenced candidate", contexts[1])
        self.assertIn("cited_by=seed-id", contexts[1])
        self.assertIn("path=seed-id -> cite-id", contexts[1])
        self.assertNotIn("Adds evidence for", contexts[1])
        self.assertTrue(contexts[1].endswith("citation text"))
        self.assertNotIn("0.2", contexts[1])

class PaperRouterTests(unittest.TestCase):
    def test_checkpoint_accepts_1024_dimension_query_and_keeps_18_labels(self) -> None:
        root = Path(__file__).resolve().parents[2]
        router = PaperNNRouter(
            root / "newbaseline/resources/router_new.pth",
            root / "newbaseline/resources/series_description.json",
        )
        selected = router.route(np.zeros(1024, dtype=np.float32), top_k=5)
        self.assertEqual(len(selected), 5)
        self.assertTrue(set(selected).issubset(PAPER_SERIES))


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
